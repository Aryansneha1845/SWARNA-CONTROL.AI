"""Generate SWARN-CONTROL.AI SANGYAN submission deck. Re-run: python generate_ppt.py"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BG = RGBColor(0x0A, 0x08, 0x06)
GOLD = RGBColor(0xF2, 0xEC, 0xE2)
AMBER = RGBColor(0xF5, 0xA3, 0x0B)
DIM = RGBColor(0x93, 0xA1, 0xB5)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
GREEN = RGBColor(0x22, 0xC5, 0x5E)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, left, top, width, height, runs, size=20, bold=False, color=GOLD, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for text, kw in runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(kw.get("size", size))
        r.font.bold = kw.get("bold", bold)
        r.font.color.rgb = kw.get("color", color)
    return box


def bullets(slide, left, top, width, height, items, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, color) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.level = 0
        r = p.add_run()
        r.text = "▸  " + text
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return box


def slide_title(sub, title, points, note=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg(s)
    add_text(s, 0.8, 0.4, 11.7, 0.6, [(sub, {"size": 15, "color": AMBER})])
    add_text(s, 0.8, 1.0, 11.7, 1.1, [(title, {"size": 38, "bold": True})])
    bullets(s, 0.8, 2.4, 11.7, 4.2, points)
    if note:
        add_text(s, 0.8, 6.6, 11.7, 0.6, [(note, {"size": 13, "color": DIM})])
    return s


# 1 — Title
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
add_text(s, 0.8, 1.2, 11.7, 0.6, [("SANGYAN 2026 · INVESTOR RESILIENCE HACKATHON · TRACK A + E", {"size": 16, "color": AMBER})])
add_text(s, 0.8, 1.9, 11.7, 1.4, [("SWARN-CONTROL.AI", {"size": 60, "bold": True})])
add_text(s, 0.8, 3.3, 11.7, 0.9, [("Secure Wealth And Resilience Network — cross-rail investigation control room.", {"size": 22, "color": DIM})])
add_text(s, 0.8, 4.3, 11.7, 0.9, [("Message + UPI + crypto trail → one evidence-backed timeline. Next control action in Hindi.", {"size": 20, "color": GOLD})])
add_text(s, 0.8, 5.6, 11.7, 0.6, [("Hum chat nahi karte, hum trail reconstruct karte hain.  ·  Solo build · Ethereum only · Zero-cost stack", {"size": 15, "color": DIM})])
add_text(s, 0.8, 6.5, 11.7, 0.5, [("Live Demo · 4-min Video · GitHub — links on last slide", {"size": 13, "color": CYAN})])

# 2 — Problem
slide_title("01 · PROBLEM", "Access outran confidence.", [
    ("16+ crore Demat accounts — 70%+ of new retail accounts from Tier-2/3 cities.", GOLD),
    ("9 out of 10 individual F&O traders incur net losses (SEBI study).", GOLD),
    ("Scam anatomy: Telegram lure → UPI transfer → crypto conversion → silence. Victim knows only when withdrawal blocks.", GOLD),
    ("Evidence is fragmented: message on Telegram, payment in UPI app, money on-chain — nobody joins the trail.", DIM),
])

# 3 — User
slide_title("02 · USER", "Praveen, 22, Ranchi.", [
    ("Hindi-first, first-generation digital-finance user. Telegram pe '2 din me double' message.", GOLD),
    ("Pays Rs 50,000 to abc-invest@upi. Coins convert, wallets split, trail goes cold.", GOLD),
    ("Existing systems fail him: SCORES forms too complex, advice too English, 'tips' too fake.", GOLD),
    ("Backup persona: Kavita — NSDL-impersonation fee scam (no crypto leg, same flow).", DIM),
], note="Solve for someone existing systems are not easy to navigate — Tier-2/3, regional language, post-scam.")

# 4 — Solution
slide_title("03 · SOLUTION", "Control room, not chatbot.", [
    ("Paste message + UPI + hash → TRACE KARO → graph: MESSAGE → UPI 12:41 → ON-RAMP 12:58 → wallets.", GOLD),
    ("Every edge clickable Evidence Card: Known fact / Supported / Possible / Unverified.", CYAN),
    ("Kill Chain (7 stages) + Hindi next-action audio + Chakshu / SCORES drafts + family alert.", GOLD),
    ("What it is NOT: no chatbox, no scam %, no tips — guardrail-safe by construction.", DIM),
])

# 5 — Journey
slide_title("04 · USER JOURNEY", "60 seconds to control.", [
    ("1. Land → What-it-does cards + 60-sec tour (first-run onboarding).", GOLD),
    ("2. ＋ New investigation → synthetic demo preset → consent tick → TRACE KARO.", GOLD),
    ("3. Read graph → play timeline → click UPI→on-ramp edge → '17-min gap + amount match = Supported'.", GOLD),
    ("4. CONTROL tab → Hindi audio plays → copy Chakshu draft → alert family → export JSON.", GOLD),
])

# 6 — Technology
slide_title("05 · TECHNOLOGY", "Deterministic money, LLM only words.", [
    ("Frontend: Next.js 14 + React Flow trail canvas + flat 2D (low-bandwidth, no 3D).", GOLD),
    ("Backend: FastAPI — /investigate, /control-pack, /tts-hindi (+ rate limits + API-key gate).", GOLD),
    ("Rules for money: resolver.py time+amount+ID joins; NetworkX-style kill chain. LLM never does money math.", CYAN),
    ("Chain: Etherscan free → Ankr no-key RPC → cached fixture fallback. Demo never dies.", GOLD),
    ("Extraction: Gemini → Groq → regex fallback. Voice: edge-tts hi-IN. OCR-ready: Tesseract.", DIM),
])

# 7 — Evidence honesty (Track E)
slide_title("06 · TRACK E · HONEST UNCERTAINTY", "No binary verdicts.", [
    ("Claim Evidence-Checker: each link labeled Known / Supported / Possible / Unverified — with WHY.", GOLD),
    ("Promotion-vs-education instinct built in: scammer's lure text framed as scammer's words, never endorsed.", GOLD),
    ("Risk language stays indicator-level (requires-review), never accusatory.", DIM),
])

# 8 — Guardrails (15% trust)
slide_title("07 · GUARDRAIL COMPLIANCE", "Public-good infrastructure.", [
    ("No tips, no price prediction, no buy/sell, no monetisation funnels — disqualification-proof.", GREEN),
    ("Privacy by design: PII redacted by default, consent-gated, in-memory only, nothing stored, no OTP/SMS harvest.", GREEN),
    ("Transparent: synthetic fixtures labeled, no fake reviews/percentages, footer + terms disclaim verdicts.", GREEN),
    ("Security audited 29 Sep: locked CORS, rate limits, security headers, Next 14.2.35, secrets server-side only.", DIM),
])

# 9 — Bharat-first (25%)
slide_title("08 · BHARAT-FIRST USABILITY", "Built for Ranchi, not just metros.", [
    ("Hindi simple-mode + hi-IN voice summary with subtitles; big fonts, 720p-friendly demo.", GOLD),
    ("Low-bandwidth: no 3D, no trackers, no third-party embeds; works on low-end devices.", GOLD),
    ("Minimal cognitive load: 60-sec tour, synthetic one-click demo, DPDP consent in plain words.", GOLD),
    ("Extensible: same flow covers IPO/phishing/F&O-tip lures; Bhashini/IVR-ready architecture.", DIM),
])

# 10 — Impact + links
slide_title("09 · IMPACT & ASK", "Hum scam nahi batate, control wapas dete hain.", [
    ("30% Resilience: victim goes from 3 disconnected fragments to 1 actionable trail + report drafts.", GOLD),
    ("15% Feasibility: zero-cost stack (free tiers, no cards) — deployable to investor-protection circles tomorrow.", GOLD),
    ("Live Demo: <vercel-url>  ·  Backend: <render-url>  ·  GitHub: github.com/Aryansneha1845/SWARNA-CONTROL.AI", CYAN),
    ("Video (4-min, Praveen case) + this deck + working prototype = complete submission.", DIM),
], note="Fill the two URLs after DEPLOY (docs/DEPLOY.md), then export this deck to PDF for submission.")

prs.save("SWARN-CONTROL.AI.pptx")
print("saved SWARN-CONTROL.AI.pptx with", len(prs.slides._sldIdLst), "slides")
